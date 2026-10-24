#!/usr/bin/env python3
"""Export Snowmoon full-book translation-review opinions to a strict 3-column XLSX.

The workbook has exactly two sheets (已采纳 / 待审核) and exactly three columns
(原翻译 / 推荐翻译 / 原因). Chapter, segment id and scope/status are carried at the start
of the 原因 cell; there are no hidden or extra data columns.

Inputs (all read-only):

  <reviews-root>/NN/opinions.json     {"chapter":N,"reviewed":true,"items":[...],"note_changes":[...]}
  <baseline-root>/NN/zh.json          immutable original Chinese ({"segments":[{"id","zh"}], "notes":[{"id","note"}]})
  <chapters-root>/NN/zh.json          editable working Chinese (same ids, same order)
  <chapters-root>/NN/<asset>          optional SVG used by scope=svg_label / svg_metadata items
  <archive-root>/assets/<file>        fixed original Chinese SVG, named by figures[].file
  <chapters-root>/NN/source.json      figure list: figures[].file/.zh map original/working assets

The archive root defaults to <repo>/inputs/upstream (the same default as
book/prepare_inputs.py) and can be overridden with --archive.  Only figures[].file
(original Chinese SVG) and figures[].zh (working copy) are used to map an asset:
filenames are never derived by guessing a "-zh" suffix.  The same mapping drives both
the per-row svg_metadata comparison and the whole-chapter asset reconciliation, and no
baseline/NN/assets copy is read or expected.

Row scopes
  body          one paragraph: zh_original must equal baseline, adopted -> working text,
                pending -> working text still baseline.
  svg_label     one visible <text>/<tspan> label, addressed by zero-based label_index.
  note          one translation note (译注) synced from the chapter's top-level
                note_changes: [{id, old, new, why[, status]}].  Every baseline->working
                note difference must be registered exactly once; a registration without a
                real difference, an unregistered difference, or a pending registration
                already written into the notes file fails.
  svg_metadata  one <title> or <desc> accessibility string of a figure asset.
                element must be title|desc and that node must be unique and keep its id
                attribute; the literal before-value must equal the archive original's
                metadata, the after-value must equal the working asset's.  Everything else in
                the asset (geometry, attributes, other visible label text) must be untouched.

Validation is deliberately all-or-nothing: the exporter refuses to write a workbook
unless every expected chapter has reviewed=true and every registered change reconciles
one-to-one with an actual baseline->working difference.  A failure exits non-zero and
writes no XLSX.

Run with:  uv run --no-project --with openpyxl python workflow/export_opinions.py --help

The exporter refuses to write a workbook unless every registered change reconciles with
the real baseline->working difference, so an unvalidated table is never produced.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
import unicodedata
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

HEADERS = ["原翻译", "推荐翻译", "原因"]
STATUSES = ("adopted", "pending")
STATUS_LABEL = {"adopted": "已采纳", "pending": "待审核"}
SHEET_NAMES = {"adopted": "已采纳", "pending": "待审核"}
REQUIRED_ITEM_KEYS = ("id", "status", "zh_original", "zh_recommended", "reason", "evidence_en")
SCOPES = ("body", "svg_label", "svg_metadata")
METADATA_ELEMENTS = ("title", "desc")
NOTE_CHANGE_KEYS = ("id", "old", "new", "why")
EVIDENCE_LIMIT = 200
COLUMN_WIDTHS = {"A": 56, "B": 56, "C": 78}
TOTAL_CHAPTERS = 32
REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ARCHIVE = REPO_ROOT / "inputs" / "upstream"


# --------------------------------------------------------------------------- #
# data model
# --------------------------------------------------------------------------- #
@dataclass
class Row:
    chapter: int
    id: str
    status: str
    zh_original: str
    zh_recommended: str
    reason: str
    evidence_en: str
    scope: str = "body"
    asset: str | None = None
    label_index: int | None = None
    element: str | None = None

    @property
    def sort_key(self) -> tuple[int, str, int, int]:
        # a metadata row sorts after the paragraph/label row sharing its base id
        scope_rank = 1 if self.scope == "svg_metadata" else 0
        return (self.chapter, self.id, scope_rank, STATUSES.index(self.status))


@dataclass
class ChapterCheck:
    number: int
    n_segments: int = 0
    n_diffs: int = 0
    adopted: int = 0
    pending: int = 0
    note_diffs: int = 0
    note_rows: int = 0
    metadata_diffs: int = 0
    metadata_rows: int = 0
    warnings: list[str] = field(default_factory=list)


@dataclass
class Validation:
    rows: list[Row]
    chapters: list[ChapterCheck]
    errors: list[str]

    @property
    def ok(self) -> bool:
        return not self.errors

    def scope_counts(self) -> dict[str, int]:
        out = {scope: 0 for scope in SCOPES}
        out["note"] = 0
        for row in self.rows:
            out[row.scope] = out.get(row.scope, 0) + 1
        return out


# --------------------------------------------------------------------------- #
# small helpers
# --------------------------------------------------------------------------- #
def read_json(path: Path, label: str, errors: list[str]):
    if not path.is_file():
        errors.append(f"{label}: 缺少文件 {path}")
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except UnicodeDecodeError as exc:
        errors.append(f"{label}: 不是 UTF-8 文本 {path}: {exc}")
    except json.JSONDecodeError as exc:
        errors.append(f"{label}: JSON 解析失败 {path}: {exc}")
    return None


def load_segments(path: Path, label: str, errors: list[str]):
    """Return (ordered ids, id->zh text) or None."""
    data = read_json(path, label, errors)
    if data is None:
        return None
    if not isinstance(data, dict):
        errors.append(f"{label}: 顶层应为对象，实际为 {type(data).__name__}")
        return None
    segs = data.get("segments")
    if not isinstance(segs, list) or not segs:
        errors.append(f"{label}: 缺少非空 segments 列表")
        return None
    order: list[str] = []
    mapping: dict[str, str] = {}
    for idx, seg in enumerate(segs):
        if not isinstance(seg, dict):
            errors.append(f"{label}: segments[{idx}] 不是对象")
            return None
        sid = seg.get("id")
        zh = seg.get("zh")
        if not isinstance(sid, str) or not sid:
            errors.append(f"{label}: segments[{idx}] 缺少字符串 id")
            return None
        if not isinstance(zh, str):
            errors.append(f"{label}: 段落 {sid} 的 zh 不是字符串")
            return None
        if sid in mapping:
            errors.append(f"{label}: 段落 id 重复 {sid}")
            return None
        order.append(sid)
        mapping[sid] = zh
    return order, mapping


def load_notes(path: Path, label: str, errors: list[str]):
    """Return {note id: literal note text} or None.  Missing/None notes -> {}."""
    data = read_json(path, label, errors)
    if data is None:
        return None
    if not isinstance(data, dict):
        errors.append(f"{label}: 顶层应为对象，实际为 {type(data).__name__}")
        return None
    notes = data.get("notes")
    if notes is None:
        return {}
    if not isinstance(notes, list):
        errors.append(f"{label}: notes 应为列表")
        return None
    out: dict[str, str] = {}
    for idx, note in enumerate(notes):
        if not isinstance(note, dict):
            errors.append(f"{label}: notes[{idx}] 不是对象")
            return None
        nid = note.get("id")
        text = note.get("note")
        if not isinstance(nid, str) or not nid:
            errors.append(f"{label}: notes[{idx}] 缺少字符串 id")
            return None
        if not isinstance(text, str):
            errors.append(f"{label}: 译注 {nid} 的 note 不是字符串")
            return None
        if nid in out:
            errors.append(f"{label}: 译注 id 重复 {nid}")
            return None
        out[nid] = text
    return out


def first_difference(a: list[str], b: list[str]) -> str:
    for idx, (x, y) in enumerate(zip(a, b)):
        if x != y:
            return f"第 {idx + 1} 项 baseline={x!r} working={y!r}"
    return f"长度不同 baseline={len(a)} working={len(b)}"


def first_sequence_difference(a: list, b: list) -> str:
    for idx, (x, y) in enumerate(zip(a, b)):
        if x != y:
            return f"第 {idx + 1} 处 baseline={x!r} working={y!r}"
    return f"元素数量不同 baseline={len(a)} working={len(b)}"


def parse_chapters(spec: str | None) -> list[int]:
    if spec is None or not spec.strip():
        return list(range(1, TOTAL_CHAPTERS + 1))
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            lo_s, hi_s = part.split("-", 1)
            lo, hi = int(lo_s), int(hi_s)
            if lo > hi:
                raise ValueError(f"章节范围倒置: {part}")
            out.extend(range(lo, hi + 1))
        else:
            out.append(int(part))
    if not out:
        raise ValueError("--chapters 为空")
    seen: set[int] = set()
    ordered: list[int] = []
    for n in sorted(out):
        if n in seen:
            continue
        seen.add(n)
        if not 1 <= n <= TOTAL_CHAPTERS:
            raise ValueError(f"章节号超出 1..{TOTAL_CHAPTERS}: {n}")
        ordered.append(n)
    return ordered


# --------------------------------------------------------------------------- #
# svg helpers
# --------------------------------------------------------------------------- #
def local_tag(node) -> str:
    return node.tag.split("}")[-1]


def load_svg(path: Path, label: str, errors: list[str]):
    try:
        return ET.parse(path).getroot()
    except (ET.ParseError, OSError) as exc:
        errors.append(f"{label}: 无法解析 SVG {path}: {exc}")
        return None


def label_texts_from_root(root) -> list[str]:
    return ["".join(node.itertext()) for node in root.iter() if local_tag(node) in ("text", "tspan")]


def svg_label_texts(path: Path, errors: list[str], label: str):
    """Literal text of every <text>/<tspan> node, in document order."""
    root = load_svg(path, label, errors)
    if root is None:
        return None
    return label_texts_from_root(root)


def metadata_nodes(root) -> dict[str, list[tuple[str | None, str]]]:
    """{title|desc: [(id attribute, literal text), ...]} in document order."""
    out: dict[str, list[tuple[str | None, str]]] = {}
    for node in root.iter():
        tag = local_tag(node)
        if tag in METADATA_ELEMENTS:
            out.setdefault(tag, []).append((node.get("id"), "".join(node.itertext())))
    return out


def svg_structure(root) -> list[tuple[str, tuple]]:
    """Document-order shape: tag + attributes of every element (text content excluded)."""
    return [(local_tag(node), tuple(sorted(node.attrib.items()))) for node in root.iter()]


def resolve_inside(chapter_dir: Path, asset, errors: list[str], label: str):
    """Return (absolute target, chapter-relative posix path) or None."""
    if not isinstance(asset, str) or not asset.strip():
        errors.append(f"{label}: 缺少 asset")
        return None
    if "\\" in asset:
        errors.append(f"{label}: asset 必须使用正斜杠相对路径: {asset!r}")
        return None
    if Path(asset).is_absolute():
        errors.append(f"{label}: asset 必须是章节目录内的相对路径: {asset!r}")
        return None
    chapter_root = chapter_dir.resolve()
    target = (chapter_dir / asset).resolve()
    try:
        rel = target.relative_to(chapter_root)
    except ValueError:
        errors.append(f"{label}: asset 越出章节目录: {asset!r}")
        return None
    if not target.is_file():
        errors.append(f"{label}: asset 文件不存在: {asset!r}")
        return None
    return target, rel.as_posix()


def figure_asset_map(chapter_dir: Path, errors: list[str]):
    """{figures[].zh working path: (figure segment id, figures[].file)}.

    The original Chinese SVG lives in <archive-root>/assets/<file>; the working copy is
    the chapter-relative path in figures[].zh.  Both must be declared in source.json, so
    neither the archive name nor the "-zh" suffix is guessed.  A figure entry that does
    not declare both is an error: an unmapped figure would silently escape the
    whole-chapter reconciliation below.
    """
    label = f"chapters/{chapter_dir.name}/source.json"
    data = read_json(chapter_dir / "source.json", label, errors)
    if data is None:
        return {}
    figures = data.get("figures")
    if not isinstance(figures, list):
        errors.append(f"{label}: 缺少 figures 列表")
        return {}
    out: dict[str, tuple[str, str]] = {}
    for idx, fig in enumerate(figures):
        if not isinstance(fig, dict):
            errors.append(f"{label}: figures[{idx}] 不是对象")
            return {}
        sid = fig.get("segment_id")
        if not isinstance(sid, str) or not sid.strip():
            errors.append(f"{label}: figures[{idx}] 缺少 segment_id 字符串")
            return {}
        zh = fig.get("zh")
        if not isinstance(zh, str) or not zh.strip():
            errors.append(
                f"{label}: figures[{idx}] 缺少工作路径 zh"
                "（原始/工作 SVG 的映射只能来自 source.json 的显式声明）"
            )
            return {}
        file = fig.get("file")
        if not isinstance(file, str) or not file.strip():
            errors.append(
                f"{label}: figures[{idx}] 缺少原始文件名 file"
                "（原始 SVG 需按 archive/assets/<file> 显式声明，不得按工作路径推测）"
            )
            return {}
        name = Path(file).name
        if name != file or ".." in Path(file).parts:
            errors.append(
                f"{label}: figures[{idx}] 的 file 必须是原始 SVG 文件名（不含目录）: {file!r}"
            )
            return {}
        prior = out.get(zh)
        if prior is not None and prior[1] != name:
            errors.append(f"{label}: figures[{idx}] 的 zh 路径 {zh!r} 重复且指向不同原始资产")
            return {}
        out[zh] = (sid, name)
    return out


def short_evidence(text: str, limit: int = EVIDENCE_LIMIT) -> str:
    flat = " ".join(str(text).split())
    if len(flat) <= limit:
        return flat
    return flat[: limit - 1].rstrip() + "…"


def reason_text(row: Row) -> str:
    head = f"【第{row.chapter}章 · {row.id}】状态：{STATUS_LABEL[row.status]}"
    if row.scope == "svg_label":
        head += f"；SVG 标签：{row.asset} 第{row.label_index} 项"
    elif row.scope == "svg_metadata":
        head += f"；SVG 无障碍文字：{row.asset} <{row.element}>"
    elif row.scope == "note":
        head += "；译注（note），源版本对齐"
    parts = [head, f"理由：{row.reason.strip()}"]
    evidence = short_evidence(row.evidence_en)
    if evidence:
        parts.append(f"英文依据：{evidence}")
    return "\n".join(parts)


# --------------------------------------------------------------------------- #
# note reconciliation
# --------------------------------------------------------------------------- #
def reconcile_notes(
    number: int,
    opinions: dict,
    base_notes: dict[str, str],
    work_notes: dict[str, str],
    errors: list[str],
) -> tuple[list[Row], int]:
    tag = f"第{number}章"
    changes = opinions.get("note_changes")
    if changes is None:
        changes = []
    if not isinstance(changes, list):
        errors.append(f"{tag}: note_changes 应为列表")
        return [], 0
    if set(base_notes) != set(work_notes):
        lost = sorted(set(base_notes) - set(work_notes))
        gained = sorted(set(work_notes) - set(base_notes))
        errors.append(f"{tag}: baseline 与工作稿的译注 id 集合不一致（缺少 {lost}，新增 {gained}），译注不得增删")
        return [], 0

    diffs = {nid for nid in base_notes if base_notes[nid] != work_notes[nid]}
    rows: list[Row] = []
    registered: dict[str, Row] = {}

    for idx, change in enumerate(changes):
        where = f"{tag} note_changes[{idx}]"
        if not isinstance(change, dict):
            errors.append(f"{where}: 不是对象")
            continue
        missing = [k for k in NOTE_CHANGE_KEYS if k not in change]
        if missing:
            errors.append(f"{where}: 缺少字段 {', '.join(missing)}")
            continue
        nid = change["id"]
        if not isinstance(nid, str) or not nid.strip():
            errors.append(f"{where}: id 必须是非空字符串")
            continue
        nid = nid.strip()
        where = f"{tag} 译注 {nid}"

        old, new, why = change["old"], change["new"], change["why"]
        bad_type = [name for name, value in (("old", old), ("new", new), ("why", why))
                    if not isinstance(value, str)]
        if bad_type:
            errors.append(f"{where}: {', '.join(bad_type)} 不是字符串")
            continue
        status = change.get("status", "adopted")
        if status not in STATUSES:
            errors.append(f"{where}: status 必须为 adopted 或 pending，实际 {status!r}")
            continue
        if not old.strip():
            errors.append(f"{where}: old 为空")
        if not new.strip():
            errors.append(f"{where}: new 为空")
        if not why.strip():
            errors.append(f"{where}: why 为空（需说明源版本依据，而不是原译错误）")
        if not (old.strip() and new.strip() and why.strip()):
            continue
        if old == new:
            errors.append(f"{where}: new 与 old 完全相同，不是真实译注改动")
            continue
        if nid in registered:
            errors.append(f"{where}: 同一译注出现多条 note_changes（每个译注只允许一行）")
            continue
        if nid not in base_notes:
            errors.append(f"{where}: baseline 中不存在该译注 id")
            continue
        if old != base_notes[nid]:
            errors.append(f"{where}: old 与 baseline 原注字面值不一致（baseline={base_notes[nid][:60]!r}…）")
        if status == "adopted":
            if work_notes[nid] == base_notes[nid]:
                errors.append(f"{where}: 登记为已同步，但工作稿译注仍是原注（未落实或属虚假记录）")
            elif work_notes[nid] != new:
                errors.append(f"{where}: 工作稿译注与 new 不一致（现有注文与登记不对应）")
        else:
            if work_notes[nid] != old:
                errors.append(f"{where}: 状态为 pending，但工作稿译注已被改动（待审核译注不得写入）")

        row = Row(
            chapter=number, id=nid, status=status,
            zh_original=old, zh_recommended=new, reason=why,
            evidence_en="", scope="note",
        )
        registered[nid] = row
        rows.append(row)

    for nid in sorted(diffs - set(registered)):
        errors.append(f"{tag} 译注 {nid}: 工作稿译注相对 baseline 有改动，但没有对应的 note_changes 记录（未登记）")
    unrealized = {nid for nid, row in registered.items() if row.status == "adopted" and nid not in diffs}
    for nid in sorted(unrealized):
        errors.append(f"{tag} 译注 {nid}: 登记为已改动，但 baseline 与工作稿译注相同（改动未落实或属虚构记录）")
    return rows, len(diffs)


# --------------------------------------------------------------------------- #
# figure asset reconciliation (title/desc metadata + untouched geometry/labels)
# --------------------------------------------------------------------------- #
def reconcile_assets(
    number: int,
    *,
    chapter_dir: Path,
    archive_root: Path,
    figures: dict[str, tuple[str, str]],
    registered: set[tuple[str, str]],
    label_cov: set[tuple[str, int]],
    errors: list[str],
) -> int:
    """Compare every working figure asset against its archive original.

    ``figures`` is the source.json mapping produced by figure_asset_map: the working
    path (figures[].zh) keys the entry, so the original named by figures[].file is the
    only comparison source.  No baseline/NN/assets copy is consulted.

    Returns the number of <title>/<desc> nodes whose literal text differs.
    """
    chapter_root = chapter_dir.resolve()
    n_diff = 0

    for work_rel, (_sid, base_name) in sorted(figures.items()):
        work_svg = chapter_dir / work_rel
        if not work_svg.is_file():
            errors.append(
                f"第{number}章 {work_rel}: source.json 已声明该工作 SVG，但文件不存在"
                "（缺图，无法核对元数据/几何/标签）"
            )
            continue
        base_svg = archive_root / "assets" / base_name
        if not base_svg.is_file():
            errors.append(
                f"第{number}章 {work_rel}: 缺少原始 SVG archive/assets/{base_name}（无法核对元数据/几何）"
            )
            continue
        broot = load_svg(base_svg, f"第{number}章 archive/assets/{base_name}", errors)
        wroot = load_svg(work_svg, f"第{number}章 chapters/{work_rel}", errors)
        if broot is None or wroot is None:
            continue
        work_rel = work_svg.resolve().relative_to(chapter_root).as_posix()

        bmeta, wmeta = metadata_nodes(broot), metadata_nodes(wroot)
        for element in METADATA_ELEMENTS:
            btexts = [text for _, text in bmeta.get(element, [])]
            wtexts = [text for _, text in wmeta.get(element, [])]
            if btexts != wtexts:
                n_diff += 1
                if (work_rel, element) not in registered:
                    errors.append(
                        f"第{number}章 {work_rel}: <{element}> 在工作稿中已与原始 SVG 不同，"
                        "但没有任何 scope=svg_metadata 意见行（未登记）"
                    )

        bstruct, wstruct = svg_structure(broot), svg_structure(wroot)
        if bstruct != wstruct:
            errors.append(
                f"第{number}章 {work_rel}: 除 title/desc 文字外的结构或属性发生变化"
                f"（可见几何不得改动；{first_sequence_difference(bstruct, wstruct)}）"
            )

        blabels, wlabels = label_texts_from_root(broot), label_texts_from_root(wroot)
        if len(blabels) != len(wlabels):
            errors.append(f"第{number}章 {work_rel}: 可见标签数量由 {len(blabels)} 变为 {len(wlabels)}")
        for index, (before, after) in enumerate(zip(blabels, wlabels)):
            if before != after and (work_rel, index) not in label_cov:
                errors.append(
                    f"第{number}章 {work_rel}: 第{index} 个可见标签由 {before!r} 变为 {after!r}，"
                    "但没有对应的 scope=svg_label 意见行（未登记标签改动）"
                )
    return n_diff


# --------------------------------------------------------------------------- #
# validation
# --------------------------------------------------------------------------- #
def validate_chapter(
    number: int,
    *,
    reviews_root: Path,
    baseline_root: Path,
    chapters_root: Path,
    archive_root: Path,
    errors: list[str],
) -> tuple[list[Row], ChapterCheck]:
    tag = f"第{number}章"
    check = ChapterCheck(number=number)
    rows: list[Row] = []

    chapter_dir = chapters_root / f"{number:02d}"
    base = load_segments(baseline_root / f"{number:02d}" / "zh.json", f"baseline/{number:02d}", errors)
    work = load_segments(chapter_dir / "zh.json", f"chapters/{number:02d}", errors)
    if base is None or work is None:
        return rows, check

    baseline_order, baseline = base
    working_order, working = work
    check.n_segments = len(working_order)

    if baseline_order != working_order:
        errors.append(
            f"{tag}: baseline 与 working 段落 id 顺序不一致（{first_difference(baseline_order, working_order)}），"
            "不能按索引对齐"
        )
        return rows, check

    diffs = {sid for sid in baseline_order if baseline[sid] != working[sid]}
    check.n_diffs = len(diffs)

    opinions = read_json(reviews_root / f"{number:02d}" / "opinions.json", f"reviews/{number:02d}/opinions.json", errors)
    if opinions is None:
        return rows, check
    if not isinstance(opinions, dict):
        errors.append(f"{tag}: opinions.json 顶层应为对象")
        return rows, check
    raw_chapter = opinions.get("chapter")
    if isinstance(raw_chapter, str) and raw_chapter.strip().isdigit():
        raw_chapter = int(raw_chapter.strip())
    if raw_chapter != number:
        errors.append(f"{tag}: opinions.json chapter={opinions.get('chapter')!r} 与目录编号不符")
    if opinions.get("reviewed") is not True:
        errors.append(f"{tag}: opinions.json reviewed 不是 true（当前 {opinions.get('reviewed')!r}），章节尚未完成审查")
        return rows, check
    items = opinions.get("items")
    if not isinstance(items, list):
        errors.append(f"{tag}: opinions.json 缺少 items 列表")
        return rows, check

    seen_ids: set[str] = set()
    seen_meta: set[tuple[str, str]] = set()
    adopted_body: set[str] = set()
    registered_meta: set[tuple[str, str]] = set()
    label_cov: set[tuple[str, int]] = set()
    # one mapping, built once: it drives both the per-row svg_metadata check and the
    # whole-chapter reconciliation, so a figure can never be checked in one place and
    # skipped in the other.
    figures = figure_asset_map(chapter_dir, errors)

    for idx, item in enumerate(items):
        where = f"{tag} items[{idx}]"
        if not isinstance(item, dict):
            errors.append(f"{where}: 不是对象")
            continue
        missing = [k for k in REQUIRED_ITEM_KEYS if k not in item]
        if missing:
            errors.append(f"{where}: 缺少字段 {', '.join(missing)}")
            continue

        sid = item["id"]
        if not isinstance(sid, str) or not sid.strip():
            errors.append(f"{where}: id 必须是非空字符串")
            continue
        sid = sid.strip()
        where = f"{tag} {sid}"

        status = item["status"]
        if status not in STATUSES:
            errors.append(f"{where}: status 必须为 adopted 或 pending，实际 {status!r}")
            continue

        original = item["zh_original"]
        recommended = item["zh_recommended"]
        reason = item["reason"]
        evidence = item["evidence_en"]
        for name, value in (("zh_original", original), ("zh_recommended", recommended),
                            ("reason", reason), ("evidence_en", evidence)):
            if not isinstance(value, str):
                errors.append(f"{where}: {name} 不是字符串")
        if not all(isinstance(v, str) for v in (original, recommended, reason, evidence)):
            continue
        if not original.strip():
            errors.append(f"{where}: zh_original 为空")
        if not recommended.strip():
            errors.append(f"{where}: zh_recommended 为空")
        if not reason.strip():
            errors.append(f"{where}: reason 为空（理由必须具体）")
        if not evidence.strip():
            errors.append(f"{where}: evidence_en 为空（需给出关键英文依据）")
        if original == recommended:
            errors.append(f"{where}: zh_recommended 与 zh_original 完全相同，不是真实修改建议")
        if not (original.strip() and recommended.strip() and reason.strip() and evidence.strip()):
            continue

        scope = item.get("scope", "body")
        if scope == "note":
            errors.append(f"{where}: 译注改动请写在 opinions.json 顶层 note_changes，items 不接受 scope=note")
            continue
        if scope not in SCOPES:
            errors.append(f"{where}: 未知 scope {scope!r}（仅支持 body / svg_label / svg_metadata）")
            continue

        if scope in ("body", "svg_label"):
            if sid in seen_ids:
                errors.append(f"{where}: 同一段落出现多行意见（应合并为一行）")
                continue
            seen_ids.add(sid)
        if sid not in working:
            errors.append(f"{where}: 工作稿中不存在该段落 id（未知 id）")
            continue
        if sid not in baseline:
            errors.append(f"{where}: baseline 中不存在该段落 id")
            continue

        row = Row(
            chapter=number, id=sid, status=status,
            zh_original=original, zh_recommended=recommended,
            reason=reason, evidence_en=evidence, scope=scope,
        )

        if scope == "body":
            if original != baseline[sid]:
                errors.append(f"{where}: zh_original 与 baseline 原译不一致（baseline={baseline[sid][:60]!r}…）")
            if status == "adopted":
                if working[sid] != recommended:
                    if working[sid] == baseline[sid]:
                        errors.append(f"{where}: 登记为已采纳，但工作稿文本未修改（漏改）")
                    else:
                        errors.append(f"{where}: 工作稿文本与 zh_recommended 不一致（工作稿改动与建议不对应）")
                adopted_body.add(sid)
            else:
                if working[sid] != original:
                    errors.append(f"{where}: 状态为 pending，但工作稿已被改动（待审核意见不得写入正文）")
            rows.append(row)

        elif scope == "svg_label":
            label_index = item.get("label_index")
            if not isinstance(label_index, int) or isinstance(label_index, bool) or label_index < 0:
                errors.append(f"{where}: scope=svg_label 需要非负整数 label_index，实际 {label_index!r}")
                continue
            resolved = resolve_inside(chapter_dir, item.get("asset"), errors, where)
            if resolved is None:
                continue
            target, rel = resolved
            texts = svg_label_texts(target, errors, where)
            if texts is None:
                continue
            if label_index >= len(texts):
                errors.append(f"{where}: label_index={label_index} 超出 {item.get('asset')} 的 text/tspan 数量 {len(texts)}")
                continue
            current = texts[label_index].strip()
            expected = recommended.strip() if status == "adopted" else original.strip()
            if current != expected:
                errors.append(
                    f"{where}: {item.get('asset')} 第{label_index} 项标签实际为 {current!r}，"
                    f"与{'推荐' if status == 'adopted' else '原译'} {expected!r} 不符"
                )
            if status == "adopted":
                label_cov.add((rel, label_index))
            row.asset = item.get("asset")
            row.label_index = label_index
            rows.append(row)

        else:  # svg_metadata
            element = item.get("element")
            if element not in METADATA_ELEMENTS:
                errors.append(f"{where}: scope=svg_metadata 的 element 必须为 title 或 desc，实际 {element!r}")
                continue
            if "label_index" in item:
                errors.append(f"{where}: scope=svg_metadata 不接受 label_index（可见标签请用 scope=svg_label）")
                continue
            asset = item.get("asset")
            resolved = resolve_inside(chapter_dir, asset, errors, where)
            if resolved is None:
                continue
            target, rel = resolved
            entry = figures.get(rel)
            if entry is None:
                errors.append(f"{where}: {asset} 未在本章 source.json figures 中声明原始资产映射（图文对应不成立）")
                continue
            if entry[0] != sid:
                errors.append(
                    f"{where}: {asset} 属于图段 {entry[0]!r}，与意见 id {sid!r} 不符（图文对应不成立）"
                )
                continue
            base_asset = archive_root / "assets" / entry[1]
            if not base_asset.is_file():
                errors.append(
                    f"{where}: 缺少原始 SVG archive/assets/{entry[1]}"
                    "（svg_metadata 必须与原始元数据逐字对照）"
                )
                continue

            work_root = load_svg(target, where, errors)
            base_root = load_svg(base_asset, where, errors)
            if work_root is None or base_root is None:
                continue
            wnodes = metadata_nodes(work_root).get(element, [])
            bnodes = metadata_nodes(base_root).get(element, [])
            if not wnodes:
                errors.append(f"{where}: 工作稿 {asset} 中不存在 <{element}> 节点")
                continue
            if len(wnodes) > 1:
                errors.append(f"{where}: 工作稿 {asset} 中 <{element}> 节点不唯一（{len(wnodes)} 个）")
                continue
            if not bnodes:
                errors.append(f"{where}: 原始 SVG 中不存在 <{element}> 节点")
                continue
            if len(bnodes) > 1:
                errors.append(f"{where}: 原始 SVG 中 <{element}> 节点不唯一（{len(bnodes)} 个）")
                continue
            work_id, work_text = wnodes[0]
            base_id, base_text = bnodes[0]
            if not base_id or not work_id:
                errors.append(f"{where}: <{element}> 必须保留非空 id 属性（baseline={base_id!r}，工作稿={work_id!r}）")
            elif work_id != base_id:
                errors.append(f"{where}: <{element}> 的 id 被改名（baseline={base_id!r}，工作稿={work_id!r}），不得改动")

            if (sid, element) in seen_meta:
                errors.append(f"{where}: 同一段落 {sid} 的 <{element}> 出现多行意见（每个节点只允许一行）")
                continue
            if (rel, element) in registered_meta:
                errors.append(f"{where}: {asset} 的 <{element}> 已有其他意见行（同一节点只允许一行）")
                continue
            seen_meta.add((sid, element))
            registered_meta.add((rel, element))

            if original != base_text:
                errors.append(
                    f"{where}: zh_original 与 baseline 资产 <{element}> 字面值不一致"
                    f"（baseline={base_text[:60]!r}…）"
                )
            if status == "adopted":
                if work_text != recommended:
                    if work_text == base_text:
                        errors.append(f"{where}: 登记为已采纳，但 {asset} 的 <{element}> 未修改（漏改）")
                    else:
                        errors.append(f"{where}: {asset} 的 <{element}> 为 {work_text[:60]!r}…，与 zh_recommended 不一致")
            else:
                if work_text != original:
                    errors.append(f"{where}: 状态为 pending，但 {asset} 的 <{element}> 已被改动（待审核元数据不得写入资源）")
            row.asset = asset
            row.element = element
            rows.append(row)

    for sid in sorted(diffs - adopted_body):
        errors.append(
            f"{tag} {sid}: 工作稿相对 baseline 有改动，但没有对应的 adopted 意见行（未登记/漏改）"
        )
    for sid in sorted(adopted_body - diffs):
        errors.append(f"{tag} {sid}: 登记为已采纳，但 baseline 与工作稿文本相同（建议未落实，或属虚构建议）")
    note_rows, note_diffs = reconcile_notes(
        number, opinions,
        load_notes(baseline_root / f"{number:02d}" / "zh.json", f"baseline/{number:02d}", errors) or {},
        load_notes(chapter_dir / "zh.json", f"chapters/{number:02d}", errors) or {},
        errors,
    )
    rows.extend(note_rows)
    check.note_diffs = note_diffs
    check.note_rows = len(note_rows)

    check.metadata_diffs = reconcile_assets(
        number,
        chapter_dir=chapter_dir,
        archive_root=archive_root,
        figures=figures,
        registered=registered_meta,
        label_cov=label_cov,
        errors=errors,
    )
    check.metadata_rows = len(registered_meta)

    check.adopted = sum(1 for r in rows if r.status == "adopted")
    check.pending = sum(1 for r in rows if r.status == "pending")
    if not rows and (diffs or note_diffs):
        errors.append(f"{tag}: 没有任何意见行，但工作稿有 {len(diffs)} 处正文改动、{note_diffs} 处译注改动")
    return rows, check


def validate(
    *,
    reviews_root: Path,
    baseline_root: Path,
    chapters_root: Path,
    archive_root: Path,
    chapters: list[int],
) -> Validation:
    errors: list[str] = []
    rows: list[Row] = []
    checks: list[ChapterCheck] = []

    for root, name in ((reviews_root, "reviews-root"), (baseline_root, "baseline-root"),
                       (chapters_root, "chapters-root")):
        if not root.is_dir():
            errors.append(f"{name}: 目录不存在 {root}")
    if errors:
        return Validation(rows=[], chapters=[], errors=errors)

    missing_chapters: set[int] = set()
    for number in chapters:
        for rel, label in ((Path("reviews") / f"{number:02d}" / "opinions.json", "reviews"),
                           (Path("baseline") / f"{number:02d}" / "zh.json", "baseline"),
                           (Path("chapters") / f"{number:02d}" / "zh.json", "chapters")):
            root = {"reviews": reviews_root, "baseline": baseline_root, "chapters": chapters_root}[label]
            if not (root / f"{number:02d}" / rel.name).is_file():
                missing_chapters.add(number)
                errors.append(f"第{number}章: 缺少 {label}/{number:02d}/{rel.name}（期望章节 {number} 尚未完成或未就绪）")

    for number in chapters:
        if number in missing_chapters:
            checks.append(ChapterCheck(number=number))
            continue
        chapter_rows, check = validate_chapter(
            number,
            reviews_root=reviews_root,
            baseline_root=baseline_root,
            chapters_root=chapters_root,
            archive_root=archive_root,
            errors=errors,
        )
        rows.extend(chapter_rows)
        checks.append(check)

    # body/svg_label share one paragraph identity; an svg_metadata row may reuse the base id.
    id_chapters: dict[str, set[int]] = {}
    for row in rows:
        id_chapters.setdefault(row.id, set()).add(row.chapter)
    for sid, chapters_seen in id_chapters.items():
        if len(chapters_seen) > 1:
            errors.append(f"段落 {sid} 在多章意见中重复出现（{sorted(chapters_seen)}）")

    rows.sort(key=lambda r: r.sort_key)
    deduped = list(dict.fromkeys(errors))
    return Validation(rows=rows, chapters=checks, errors=deduped)


# --------------------------------------------------------------------------- #
# workbook
# --------------------------------------------------------------------------- #
def _set_text(cell, text: str) -> None:
    cell.value = text
    # openpyxl marks strings starting with "=" as formulas; force literal text.
    cell.data_type = "s"


def _display_width(text: str) -> int:
    total = 0
    for ch in text:
        total += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return total


def _line_count(text: str, width: int) -> int:
    lines = 0
    for part in text.split("\n"):
        lines += max(1, math.ceil(_display_width(part) / max(width - 2, 8)))
    return lines


def write_workbook(rows: list[Row], output: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F3864")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for status in STATUSES:
        ws = wb.create_sheet(SHEET_NAMES[status])
        for col, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col)
            _set_text(cell, header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[1].height = 22
        for column, width in COLUMN_WIDTHS.items():
            ws.column_dimensions[column].width = width

        selected = [r for r in rows if r.status == status]
        for offset, row in enumerate(selected, start=2):
            texts = [row.zh_original, row.zh_recommended, reason_text(row)]
            for col, text in enumerate(texts, start=1):
                cell = ws.cell(row=offset, column=col)
                _set_text(cell, text)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
            lines = max(
                _line_count(texts[0], COLUMN_WIDTHS["A"]),
                _line_count(texts[1], COLUMN_WIDTHS["B"]),
                _line_count(texts[2], COLUMN_WIDTHS["C"]),
            )
            ws.row_dimensions[offset].height = min(15 * lines + 4, 409)

        last_row = max(2, len(selected) + 1)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:C{last_row}"
        ws.sheet_view.showGridLines = True

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def write_markdown(rows: list[Row], path: Path, title: str) -> None:
    def esc(text: str) -> str:
        return text.replace("|", "\\|").replace("\r\n", "\n").replace("\n", "<br>")

    lines = [f"# {title}", "", "| 原翻译 | 推荐翻译 | 原因 |", "| --- | --- | --- |"]
    for row in rows:
        lines.append(f"| {esc(row.zh_original)} | {esc(row.zh_recommended)} | {esc(reason_text(row))} |")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# --------------------------------------------------------------------------- #
# cli
# --------------------------------------------------------------------------- #
def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="把 reviews/NN/opinions.json 导出为严格三列 Excel（已采纳 / 待审核）。",
        epilog="校验不通过时退出码为 3 且不写出工作簿；参数错误为 2。",
    )
    p.add_argument("--reviews-root", required=True, type=Path, help="reviews 根目录（含 NN/opinions.json）")
    p.add_argument("--baseline-root", required=True, type=Path, help="baseline 根目录（含 NN/zh.json 原译）")
    p.add_argument("--chapters-root", required=True, type=Path, help="chapters 根目录（含 NN/zh.json 工作稿）")
    p.add_argument("--output", type=Path, help="输出 .xlsx 路径")
    p.add_argument("--markdown", type=Path, help="可选：同时输出三列 Markdown 副本")
    p.add_argument("--report", type=Path, help="可选：写出 JSON 校验报告")
    p.add_argument("--chapters", help="期望章节，如 1,5,9 或 1-32；默认 1..32")
    p.add_argument("--archive", type=Path,
                   help="原始中文 SVG 归档根目录（默认 <repo>/inputs/upstream，与 book/prepare_inputs.py 对齐）")
    p.add_argument("--validate-only", action="store_true", help="只校验，不写工作簿")
    p.add_argument("--title", default="Snowmoon 翻译修改意见（全书）", help="Markdown 标题")
    return p


def report_payload(validation: Validation, args: argparse.Namespace, chapters: list[int],
                   output: Path | None, sha256: str | None) -> dict:
    scope_counts = validation.scope_counts()
    return {
        "tool": "workflow/export_opinions.py",
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "status": "ok" if validation.ok else "failed",
        "expected_chapters": chapters,
        "reviews_root": str(args.reviews_root),
        "baseline_root": str(args.baseline_root),
        "chapters_root": str(args.chapters_root),
        "archive_root": str(args.archive_root),
        "output": str(output) if output else None,
        "output_sha256": sha256,
        "counts": {
            "chapters": len(validation.chapters),
            "segments": sum(c.n_segments for c in validation.chapters),
            "baseline_working_diffs": sum(c.n_diffs for c in validation.chapters),
            "adopted": sum(1 for r in validation.rows if r.status == "adopted"),
            "pending": sum(1 for r in validation.rows if r.status == "pending"),
            "rows_total": len(validation.rows),
            "rows_by_scope": scope_counts,
            "body_rows": scope_counts.get("body", 0),
            "note_diffs": sum(c.note_diffs for c in validation.chapters),
            "note_changes": scope_counts.get("note", 0),
            "svg_label_changes": scope_counts.get("svg_label", 0),
            "svg_metadata_diffs": sum(c.metadata_diffs for c in validation.chapters),
            "svg_metadata_changes": scope_counts.get("svg_metadata", 0),
        },
        "per_chapter": [
            {
                "chapter": c.number,
                "segments": c.n_segments,
                "baseline_working_diffs": c.n_diffs,
                "adopted": c.adopted,
                "pending": c.pending,
                "note_diffs": c.note_diffs,
                "note_rows": c.note_rows,
                "svg_metadata_diffs": c.metadata_diffs,
                "svg_metadata_rows": c.metadata_rows,
                "warnings": c.warnings,
            }
            for c in validation.chapters
        ],
        "errors": validation.errors,
    }


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        chapters = parse_chapters(args.chapters)
    except ValueError as exc:
        parser.error(f"--chapters: {exc}")
        return 2
    if not args.validate_only and args.output is None:
        parser.error("需要 --output（或用 --validate-only 只做校验）")
        return 2
    args.archive_root = args.archive if args.archive is not None else DEFAULT_ARCHIVE

    validation = validate(
        reviews_root=args.reviews_root,
        baseline_root=args.baseline_root,
        chapters_root=args.chapters_root,
        archive_root=args.archive_root,
        chapters=chapters,
    )

    output = None if args.validate_only else args.output
    sha = None
    if validation.ok and output is not None:
        write_workbook(validation.rows, output)
        sha = hashlib.sha256(output.read_bytes()).hexdigest()
        if args.markdown:
            write_markdown(validation.rows, args.markdown, args.title)

    if args.report:
        payload = report_payload(validation, args, chapters, output, sha)
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    counts = {}
    for row in validation.rows:
        counts[row.status] = counts.get(row.status, 0) + 1
    diffs = sum(c.n_diffs for c in validation.chapters)
    if not validation.ok:
        print(f"校验失败：{len(validation.errors)} 个问题，未写出工作簿。", file=sys.stderr)
        for err in validation.errors:
            print(f"  - {err}", file=sys.stderr)
        if output is not None and output.exists():
            print(f"注意：{output} 已存在，是此前运行结果，不可当作本次结果使用。", file=sys.stderr)
        return 3

    scopes = validation.scope_counts()
    print(
        f"OK 章节 {len(validation.chapters)}/{len(chapters)} · 段落 {sum(c.n_segments for c in validation.chapters)} · "
        f"工作稿改动 {diffs} · 已采纳 {counts.get('adopted', 0)} · 待审核 {counts.get('pending', 0)} · "
        f"译注 {scopes.get('note', 0)} · SVG标签 {scopes.get('svg_label', 0)} · "
        f"SVG元数据 {scopes.get('svg_metadata', 0)} · 行合计 {len(validation.rows)}"
    )
    if output is not None:
        print(f"工作簿：{output}  sha256={sha}")
    if args.report:
        print(f"报告：{args.report}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
